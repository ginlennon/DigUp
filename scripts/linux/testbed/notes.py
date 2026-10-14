"""Notes in the formats macOS's NSAttributedString writes: .rtf (as TextEdit/Cocoa writes it: non-ASCII as \\uN escapes)
and .docx (one paragraph per line, python-docx); .odt as a minimal OpenDocument text (one paragraph per line); anything
else is plain UTF-8, as `String.write(to:)` writes it."""
import datetime
import zipfile
from xml.sax.saxutils import escape as xml_escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

FIXED_TIME = datetime.datetime(2026, 10, 1, 12, 0, 0)


def rtf(content, font="Helvetica", size=13, rtl=False):
    def escape(ch):
        if ch in "\\{}":
            return "\\" + ch
        if ch == "\n":
            return "\\\n"  # Cocoa ends a paragraph with a backslash and a newline
        if ord(ch) < 128:
            return ch
        if ord(ch) > 0xFFFF:  # outside the BMP: a UTF-16 surrogate pair, each as a signed 16-bit \uN
            v = ord(ch) - 0x10000
            return "".join(f"\\u{u - 65536 if u > 32767 else u} " for u in (0xD800 + (v >> 10), 0xDC00 + (v & 0x3FF)))
        return f"\\u{ord(ch) - 65536 if ord(ch) > 32767 else ord(ch)} "  # the space ends the control word

    paragraph = "\\pard\\tx560\\tx1120\\tx1680\\tx2240\\tx2800\\tx3360\\tx3920\\tx4480\\tx5040\\tx5600\\tx6160\\tx6720"
    paragraph += "\\rtlpar\\qr\\partightenfactor0" if rtl else "\\pardirnatural\\partightenfactor0"
    return ("{\\rtf1\\ansi\\ansicpg1252\\cocoartf2822\n"
            "\\cocoatextscaling0\\cocoaplatform0{\\fonttbl\\f0\\fnil\\fcharset0 " + font + ";}\n"
            "{\\colortbl;\\red255\\green255\\blue255;}\n"
            "{\\*\\expandedcolortbl;;}\n"
            "\\paperw11900\\paperh16840\\margl1440\\margr1440\\vieww11520\\viewh8400\\viewkind0\n"
            + paragraph + "\n\n"
            f"\\f0\\fs{size * 2} \\cf0 \\uc0 " + "".join(escape(ch) for ch in content) + "}")


def docx(path, content, font="Helvetica", size=13, rtl=False):
    document = Document()
    document.core_properties.created = FIXED_TIME
    document.core_properties.modified = FIXED_TIME
    document.core_properties.author = "DigUp testbed"
    for line in content.split("\n"):
        paragraph = document.add_paragraph()
        if rtl:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph._p.get_or_add_pPr().append(paragraph._p.makeelement(qn("w:bidi"), {}))
        if line:
            run = paragraph.add_run(line)
            run.font.size = Pt(size)
            run.font.name = font
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font)
    document.save(path)


def odt(path, content):
    """The smallest OpenDocument text LibreOffice opens: the mimetype first and stored, a manifest, content.xml."""
    office = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
    body = "".join(f"<text:p>{xml_escape(line)}</text:p>" if line else "<text:p/>" for line in content.split("\n"))
    xml = ('<?xml version="1.0" encoding="UTF-8"?>\n'
           f'<office:document-content xmlns:office="{office}" '
           'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.3">'
           f"<office:body><office:text>{body}</office:text></office:body></office:document-content>")
    manifest = ('<?xml version="1.0" encoding="UTF-8"?>\n'
                '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
                'manifest:version="1.3">'
                '<manifest:file-entry manifest:full-path="/" manifest:version="1.3" '
                'manifest:media-type="application/vnd.oasis.opendocument.text"/>'
                '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
                "</manifest:manifest>")
    stamp = FIXED_TIME.timetuple()[:6]
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype", stamp), "application/vnd.oasis.opendocument.text",
                   compress_type=zipfile.ZIP_STORED)
        for name, data in [("META-INF/manifest.xml", manifest), ("content.xml", xml)]:
            z.writestr(zipfile.ZipInfo(name, stamp), data, compress_type=zipfile.ZIP_DEFLATED)


def write(path, content, font="Helvetica", size=13, rtl=False):
    if path.endswith(".odt"):
        odt(path, content)
    elif path.endswith(".rtf"):
        with open(path, "w", encoding="ascii") as f:
            f.write(rtf(content, font, size, rtl))
    elif path.endswith(".docx"):
        docx(path, content, font, size, rtl)
    else:
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
